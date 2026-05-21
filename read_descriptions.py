import openpyxl
import sys

# Set UTF-8 encoding for output
sys.stdout.reconfigure(encoding='utf-8')

# Load the workbook
wb = openpyxl.load_workbook('excel_db/assets_raw_100526.xlsx')
ws = wb.active

# Get the header row
header = [cell.value for cell in ws[1]]
print('Header row:')
for i, h in enumerate(header):
    print(f'  Column {i}: {h}')
print()

# Extract first 20 rows from column D
print('First 20 rows of Column D (Description):')
print('=' * 100)
for row_num in range(2, 22):  # Rows 2-21 (row 1 is header)
    cell_value = ws[f'D{row_num}'].value
    row_display = row_num - 2  # Display as 0-19
    if cell_value is not None:
        print(f'Row {row_display}: {str(cell_value)}')
    else:
        print(f'Row {row_display}: [Empty]')
